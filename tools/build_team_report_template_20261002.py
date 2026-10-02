"""Build an editable team-report layout for native Hancom HWP conversion."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'exports' / 'reports' / 'team-midterm-template-20261002'
OUT.mkdir(parents=True, exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.75), Inches(.7)
sec.left_margin, sec.right_margin = Inches(.8), Inches(.8)
sec.header_distance = sec.footer_distance = Inches(.3)

def font(style, size, bold=False):
    style.font.name = '맑은 고딕'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '맑은 고딕')

font(doc.styles['Normal'], 11)
doc.styles['Normal'].paragraph_format.line_spacing = 1.25
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
for name, size in [('Title', 24), ('Subtitle', 14), ('Heading 1', 18), ('Heading 2', 12)]:
    font(doc.styles[name], size, name != 'Subtitle')
    st = doc.styles[name]
    st.paragraph_format.space_before = Pt(10 if name == 'Heading 2' else 0)
    st.paragraph_format.space_after = Pt(7)
    st.paragraph_format.keep_with_next = True
    for borders in list(st.element.findall('.//' + qn('w:pBdr'))):
        borders.getparent().remove(borders)

header = sec.header.paragraphs[0]
header.text = '버스태워줘  |  2026년 2학기 자율설계학기제 중간 팀보고서'
header.paragraph_format.space_after = Pt(0)
for run in header.runs:
    run.font.size = Pt(9)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

def paragraph(text='', small=False, after=6):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(after)
    if small:
        for run in p.runs:
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor.from_string('595959')
    return p

def heading(text, level=2):
    return doc.add_heading(text, level)

def slot(title, guidance):
    heading(title)
    paragraph(guidance, small=True)
    p = paragraph('[본문 작성]', after=8)
    for run in p.runs:
        run.font.color.rgb = RGBColor.from_string('7F7F7F')

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge)
        for key, val in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            e.set(qn('w:' + key), val)
        borders.append(e)
    pr.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for edge in ('top', 'bottom', 'left', 'right'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:w'), '90' if edge in ('top', 'bottom') else '110')
        e.set(qn('w:type'), 'dxa')
        margins.append(e)
    pr.append(margins)
    for i, width in enumerate(widths):
        t.columns[i].width = Inches(width)
    for i, text in enumerate(headers):
        t.rows[0].cells[i].text = text
        shade = OxmlElement('w:shd')
        shade.set(qn('w:fill'), 'E7E6E6')
        t.rows[0].cells[i]._tc.get_or_add_tcPr().append(shade)
    repeat = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for texts in rows:
        cells = t.add_row().cells
        for i, text in enumerate(texts):
            cells[i].text = text
    for row_idx, row in enumerate(t.rows):
        no_split = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(no_split)
        for col, cell in enumerate(row.cells):
            cell.width = Inches(widths[col])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                if col == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(10)
                    run.font.bold = row_idx == 0
    paragraph('', after=3)
    return t

def figure(number, text):
    paragraph(f'[그림 {number} 삽입 위치] {text}', small=True, after=5)
    p = paragraph('사진 또는 화면 캡처를 넣고 아래에 설명을 작성한다.', small=True)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(14)
    paragraph(f'그림 {number}. [그림 설명 작성]', small=True)

chapters = [
    '프로젝트 개요',
    '공통 실내지도 제작 및 활용',
    '셔틀버스 위치 및 도착정보 시스템',
    '정류장 대기 인원 및 이용 안내',
    '실내측위 및 강의실 길찾기',
    '앱 기능 및 서버 연계',
    '중간 성과 및 향후 계획',
]

paragraph('', after=40)
p = doc.add_paragraph('2026년 2학기 자율설계학기제', style='Subtitle')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph('중간 팀보고서', style='Title')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
paragraph('', after=18)
p = paragraph('셔틀버스 이용 정보와 교내 실내 안내 서비스')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
paragraph('', after=28)
paragraph('팀명: 버스태워줘')
paragraph('과제명: [계획서의 공식 과제명 입력]')
paragraph('보고 기간: [실제 보고 기간 입력]')
paragraph('지도교수: [성명 입력]')
paragraph('제출일: [제출일 입력]')
paragraph('', after=10)
table(['팀원', '학번', '담당 분야'], [
    ['김태진', '20222963', '[담당 분야 작성]'],
    ['김정민', '20222926', '[담당 분야 작성]'],
    ['오한울', '20222992', '[담당 분야 작성]'],
    ['정재모', '20222994', '[담당 분야 작성]'],
    ['진현호', '20222967', '[담당 분야 작성]'],
], [1.0, 1.35, 4.55])

doc.add_page_break()
heading('목차', 1)
for i, name in enumerate(chapters, 1):
    p = paragraph(f'{i}  {name}', after=12)
    p.paragraph_format.space_before = Pt(5)
paragraph('부록  회의 기록과 참고자료', after=22)
heading('작성 방법')
paragraph('각 장은 기능별로 구성하였다. 담당자는 개발이 필요한 이유, 구현 방법, 시행착오와 수정, 확인한 결과와 남은 문제를 이어서 작성한다.')
paragraph('대괄호로 표시한 안내와 빈칸을 실제 내용으로 바꾼다. 표에는 같은 조건으로 비교한 결과를 넣고, 사진에는 무엇을 확인할 수 있는지 설명한다.')
paragraph('구현 완료, 개발 중, 계획을 구분한다. 지도 활용 내용은 2장에, 분실물 게시판은 6장에 작성한다.', small=True)

doc.add_page_break()
heading('1 프로젝트 개요', 1)
slot('1 1 추진 배경과 기존 서비스의 한계', '기존 셔틀버스 정보 제공의 한계와 이용자가 겪는 불편을 설명한다.')
slot('1 2 목표와 서비스 구성', '버스 위치와 도착정보, 정류장 이용 안내, 실내 강의실 안내가 어떻게 연결되는지 작성한다.')
slot('1 3 차별화 방향과 주제 보완', '교수 자문을 바탕으로 검토한 의견과 실제 채택한 기능, 선택 이유를 작성한다.')
slot('1 4 팀원별 역할과 협업', '계획서의 업무분장과 현재 실제 담당 업무, 공통 지도와 서버를 공유한 방식을 설명한다.')
figure('1', '전체 서비스 구성도')

doc.add_page_break()
heading('2 공통 실내지도 제작 및 활용', 1)
slot('2 1 제작 목적과 기초 자료', '실내측위와 강의실 상태 확인에 공통 지도가 필요한 이유, 피난안내도와 현장 자료의 사용 방법을 설명한다.')
slot('2 2 층별 2D 지도 제작', '공간 구분, 강의실 번호, 복도와 계단의 연결을 지도에 옮긴 과정과 현장 수정 내용을 작성한다.')
slot('2 3 3D 지도 제작', '확인한 2D 지도를 바탕으로 벽과 공간을 구성한 방법, 3D 촬영 자료와 대조한 내용을 작성한다.')
slot('2 4 실제 거리와 지도 정합', '실제 통행 구간을 어떻게 측정했으며, 지도 표시와 위치 계산의 거리 기준을 어떻게 맞췄는지 작성한다.')
slot('2 5 시간표 기반 강의실 상태 확인', '강의실 지도와 시간표를 연결해 수업 여부와 이용 가능한 시간대를 보여주는 방법을 작성한다.')
figure('2', '같은 층의 2D와 3D 지도 및 강의실 상태 화면')

doc.add_page_break()
heading('3 셔틀버스 위치 및 도착정보 시스템', 1)
slot('3 1 시스템 목적과 데이터 흐름', '차량 위치 측정, 무선 전달, 서버 저장, 이용자 화면 표시의 흐름을 설명한다.')
slot('3 2 차량 위치 수집과 통신', '사용 장치, GPS 위치 수집, LoRa 무선 전달 방법과 설정을 선택한 이유를 작성한다.')
slot('3 3 서버 연계와 도착정보 제공', '수신한 정보를 지도에 표시하는 방법, 갱신 주기, 도착정보 계산과 처리 방식을 작성한다.')
slot('3 4 통신 시험과 개선 결과', '통신이 잘되지 않았던 조건과 변경 사항, 시험 후 확인한 결과를 작성한다.')
table(['시험 항목', '시험 조건', '확인 결과', '남은 문제'], [
    ['위치 갱신', '[조건]', '[결과]', '[문제]'],
    ['통신 거리와 안정성', '[조건]', '[결과]', '[문제]'],
], [1.3, 1.8, 1.9, 1.9])
figure('3', '장치 연결 구성 또는 버스 위치 화면')

doc.add_page_break()
heading('4 정류장 대기 인원 및 이용 안내', 1)
slot('4 1 기능 목적과 관측 구역', '도착정보만으로 알기 어려운 대기 인원과 혼잡 상황, 카메라 관측 구역을 설명한다.')
slot('4 2 카메라 수집과 사람 인식', '영상 수집과 전달, 사람 검출, 같은 사람을 이어서 추적하는 방법을 작성한다.')
slot('4 3 대기 인원과 혼잡 안내', '노선별 대기 인원 구분, 서버 반영, 화면과 음성 안내의 동작 과정을 작성한다.')
slot('4 4 실험 결과와 추가 기능 검토', '가림과 중복 집계 등 시행착오, 개선 결과를 작성하고 버스 승하차 인원과 탑승 가능 안내의 현재 진행 상태를 구분한다.')
table(['관측 상황', '실제 인원', '집계 결과', '오류와 개선'], [
    ['[시험 상황 1]', '[명]', '[명]', '[원인과 수정]'],
    ['[시험 상황 2]', '[명]', '[명]', '[원인과 수정]'],
], [1.55, 1.1, 1.1, 3.15])
figure('4', '사람 인식 화면과 정류장 안내 화면')

doc.add_page_break()
heading('5 실내측위 및 강의실 길찾기', 1)
slot('5 1 실내 위치 추정과 Android 앱', 'GPS를 쓰기 어려운 실내에서 휴대폰의 걸음과 방향을 누적하는 PDR, 앱의 위치 표시와 수집 기능을 설명한다.')
slot('5 2 지도 제약과 위치 후보 관리', '복도 밖 이동을 제한한 초기 방식과 2층 및 3층 추가 영역에서 여러 위치 후보를 유지하도록 바꾼 이유와 방법을 작성한다.')
slot('5 3 경로 라벨과 센서 기준점 수집', '이동 중 통과 지점과 센서를 기록하는 경로 라벨 수집, 정지 상태의 장소별 신호를 모으는 기준점 수집을 구분해 설명한다.')
slot('5 4 보정 방식과 엔진 비교', '보행 중 자기장과 BLE, 기준점의 Wi-Fi 신호를 위치 후보에 반영하는 방법과 보폭 모델 비교 내용을 작성한다.')
slot('5 5 평가 결과와 남은 오차', '사용 자료 수와 평가 조건, 개선된 부분과 해결되지 않은 오차를 작성한다.')
table(['비교 구성', '자료와 평가 조건', '결과 및 판단'], [
    ['기본 방식', '[자료와 조건]', '[결과]'],
    ['보정 추가 방식', '[동일 조건]', '[개선과 한계]'],
], [1.4, 2.6, 2.9])
figure('5', '앱의 예상 위치와 실제 통과 지점 비교')

doc.add_page_break()
heading('6 앱 기능 및 서버 연계', 1)
slot('6 1 이용 화면과 기능 구성', '셔틀버스 정보, 정류장 안내, 실내지도와 학생 편의 기능의 화면 구성을 작성한다.')
slot('6 2 분실물 게시판', '분실물 정보와 사진을 등록하고 목록 및 상세 화면에서 확인하는 과정, 저장 방식을 설명한다.')
slot('6 3 공통 데이터와 서버 연결', '차량 위치와 인원 정보 등 어떤 데이터를 주고받으며, Firebase와 공통 지도 데이터를 어떻게 연결하는지 작성한다.')
slot('6 4 모바일 적용과 연계 검사', '휴대폰 화면에 맞춘 구성, 앱 실행 및 서버 정보 표시를 확인한 내용과 남은 문제를 작성한다.')
table(['기능', '연결 정보', '현재 상태', '확인 방법'], [
    ['버스와 정류장 안내', '[정보]', '[상태]', '[검사]'],
    ['지도와 강의실 상태', '[정보]', '[상태]', '[검사]'],
    ['분실물 게시판', '[정보]', '[상태]', '[검사]'],
], [1.55, 1.9, 1.15, 2.3])
figure('6', '앱 주요 화면과 분실물 게시판')

doc.add_page_break()
heading('7 중간 성과 및 향후 계획', 1)
slot('7 1 주요 결과물과 목표 달성 현황', '앞 장에서 확인한 지도, 장치, 앱, 실험 자료를 요약하고 실제 완료와 진행 중인 항목을 구분한다.')
table(['분야', '주요 결과물', '현재 상태', '다음 확인'], [
    ['공통 지도와 활용', '[결과물]', '[상태]', '[확인할 항목]'],
    ['셔틀버스 정보', '[결과물]', '[상태]', '[확인할 항목]'],
    ['정류장 이용 안내', '[결과물]', '[상태]', '[확인할 항목]'],
    ['실내측위', '[결과물]', '[상태]', '[확인할 항목]'],
    ['앱과 서버', '[결과물]', '[상태]', '[확인할 항목]'],
], [1.4, 2.15, 1.1, 2.25])
slot('7 2 향후 일정과 검증 계획', '실제 셔틀 운행 시험, 정류장 관측, 설치 비콘을 이용한 실내 보정 시험 등 우선 과제와 담당자를 작성한다.')
table(['과제', '담당', '목표 시기', '완료 확인 기준'], [
    ['[우선 과제 1]', '[담당자]', '[시기]', '[확인 기준]'],
    ['[우선 과제 2]', '[담당자]', '[시기]', '[확인 기준]'],
], [2.15, 1.0, 1.05, 2.7])
slot('7 3 전체 서비스 연계', '실외 버스 이동 정보부터 건물 진입 후 실내 강의실 안내까지 연결하는 단계와 시험 시나리오를 작성한다.')

doc.add_page_break()
heading('부록 회의 기록과 참고자료', 1)
heading('회의 기록 양식')
paragraph('실제로 진행한 회의의 날짜, 참석자, 의견과 결정 사항을 작성한다.', small=True)
paragraph('회의 일시: [입력]     장소: [입력]')
paragraph('참석자: [입력]')
paragraph('안건: [입력]')
paragraph('주요 의견: [입력]')
paragraph('결정 사항과 후속 작업: [입력]')
paragraph('', after=15)
heading('참고자료와 결과물 목록')
paragraph('피난안내도, 현장 측정 자료, 팀원 개인보고서, 장치 문서, 실험 기록의 출처와 보관 위치를 작성한다.', small=True)
table(['구분', '자료 또는 결과물', '출처와 보관 위치'], [
    ['지도', '[자료명]', '[출처 또는 경로]'],
    ['실험', '[자료명]', '[날짜와 경로]'],
    ['개발 결과', '[파일 또는 저장소]', '[위치와 버전]'],
    ['참고 문헌', '[제목]', '[저자와 링크]'],
], [1.2, 2.5, 3.2])

doc.core_properties.title = '2026년 2학기 자율설계학기제 중간 팀보고서 기본틀'
doc.core_properties.author = '버스태워줘'
doc.core_properties.subject = '지도 제작과 활용 셔틀버스 정보 정류장 안내 실내측위 앱 기능'
dest = OUT / 'team-report-template.docx'
doc.save(dest)
(OUT / 'template-spec.json').write_text(json.dumps({
    'purpose': '편집 가능한 팀 중간보고서 기본틀',
    'chapters': chapters,
    'excluded': ['목적지를 선택하면 교내 이동경로 표시', '단과대학별 빠른 길찾기와 지름길 안내'],
    'target_format': 'HWP',
    'source_docx': str(dest),
    'intended_pages': 10,
}, ensure_ascii=False, indent=2), encoding='utf-8')
print(dest)
